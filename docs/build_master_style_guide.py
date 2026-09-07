from pathlib import Path
import re

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path("/Users/fernandoliernur/Projects/nutrisoft")
SOURCE = ROOT / "docs" / "NUTRISOFT_MASTER_STYLE_AND_REQUIREMENTS.md"
OUTPUT = ROOT / "docs" / "NutriSoft_Documento_Maestro_Estilo_y_Requisitos_2026-08-14.docx"

INK = "151B22"
SECONDARY = "66727D"
TERTIARY = "8A959D"
BRAND = "55AEB8"
BRAND_STRONG = "357984"
BRAND_SOFT = "DDF3F2"
BORDER = "E2E9EC"
SURFACE_SUBTLE = "F2F7F8"
INFO_BG = "EAEFFC"
WARNING_BG = "FDF6E2"
CRITICAL_BG = "FCEBEA"
SUCCESS_BG = "E8F5EE"
TABLE_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120


def set_font(run, name="Calibri", size=None, color=None, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, bottom=80, start=120, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("bottom", bottom), ("start", start), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def prevent_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:cantSplit"))


def set_table_geometry(table, widths_dxa):
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for cell, width in zip(row.cells, widths_dxa):
            cell.width = Inches(width / 1440)
            tc_w = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                cell._tc.get_or_add_tcPr().append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("NutriSoft · Documento maestro  |  ")
    set_font(run, size=8, color=TERTIARY)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, end])


def add_numbering(doc):
    numbering = doc.part.numbering_part.element

    def abstract(abstract_id, num_fmt, text, left, hanging, font=None):
        abs_num = OxmlElement("w:abstractNum")
        abs_num.set(qn("w:abstractNumId"), str(abstract_id))
        multi = OxmlElement("w:multiLevelType")
        multi.set(qn("w:val"), "singleLevel")
        abs_num.append(multi)
        lvl = OxmlElement("w:lvl")
        lvl.set(qn("w:ilvl"), "0")
        start = OxmlElement("w:start")
        start.set(qn("w:val"), "1")
        fmt = OxmlElement("w:numFmt")
        fmt.set(qn("w:val"), num_fmt)
        lvl_text = OxmlElement("w:lvlText")
        lvl_text.set(qn("w:val"), text)
        suff = OxmlElement("w:suff")
        suff.set(qn("w:val"), "tab")
        ppr = OxmlElement("w:pPr")
        tabs = OxmlElement("w:tabs")
        tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), "num")
        tab.set(qn("w:pos"), str(left))
        tabs.append(tab)
        ind = OxmlElement("w:ind")
        ind.set(qn("w:left"), str(left))
        ind.set(qn("w:hanging"), str(hanging))
        spacing = OxmlElement("w:spacing")
        spacing.set(qn("w:after"), "80")
        spacing.set(qn("w:line"), "300")
        spacing.set(qn("w:lineRule"), "auto")
        ppr.extend([tabs, ind, spacing])
        lvl.extend([start, fmt, lvl_text, suff, ppr])
        if font:
            rpr = OxmlElement("w:rPr")
            rfonts = OxmlElement("w:rFonts")
            rfonts.set(qn("w:ascii"), font)
            rfonts.set(qn("w:hAnsi"), font)
            rpr.append(rfonts)
            lvl.append(rpr)
        abs_num.append(lvl)
        numbering.append(abs_num)
        num = OxmlElement("w:num")
        num_id = abstract_id + 1
        num.set(qn("w:numId"), str(num_id))
        abs_ref = OxmlElement("w:abstractNumId")
        abs_ref.set(qn("w:val"), str(abstract_id))
        num.append(abs_ref)
        numbering.append(num)
        return num_id

    return abstract(50, "bullet", "•", 540, 270, "Calibri"), abstract(60, "decimal", "%1.", 540, 270)


def add_numbering_instance(doc, abstract_id, num_id):
    """Create an independent numbering sequence backed by an existing definition."""
    numbering = doc.part.numbering_part.element
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abs_ref = OxmlElement("w:abstractNumId")
    abs_ref.set(qn("w:val"), str(abstract_id))
    num.append(abs_ref)
    level_override = OxmlElement("w:lvlOverride")
    level_override.set(qn("w:ilvl"), "0")
    start_override = OxmlElement("w:startOverride")
    start_override.set(qn("w:val"), "1")
    level_override.append(start_override)
    num.append(level_override)
    numbering.append(num)
    return num_id


def set_num(paragraph, num_id):
    ppr = paragraph._p.get_or_add_pPr()
    num_pr = ppr.get_or_add_numPr()
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    numid = OxmlElement("w:numId")
    numid.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, numid])


def add_inline_runs(paragraph, text, size=11, color=INK):
    parts = re.split(r"(`[^`]+`|\*\*[^*]+\*\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            set_font(run, size=size, color=color, bold=True)
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            set_font(run, name="Consolas", size=max(8.5, size - 1), color=BRAND_STRONG)
            rpr = run._element.get_or_add_rPr()
            shd = OxmlElement("w:shd")
            shd.set(qn("w:fill"), SURFACE_SUBTLE)
            rpr.append(shd)
        else:
            # LibreOffice can collapse regular spaces next to shaded inline-code
            # runs. Non-breaking boundary spaces preserve readable separation in
            # both DOCX and rendered PDF output.
            safe_part = part
            if safe_part.startswith(" "):
                safe_part = "\u00a0" + safe_part[1:]
            if safe_part.endswith(" "):
                safe_part = safe_part[:-1] + "\u00a0"
            run = paragraph.add_run(safe_part)
            set_font(run, size=size, color=color)


def add_section_heading(doc, text, level):
    if text == "4.3 Esquemas de backend":
        doc.add_page_break()
    # A short explicit spacer prevents LibreOffice from visually collapsing a
    # heading onto the final line of a preceding numbered/bulleted paragraph.
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_before = Pt(0)
    spacer.paragraph_format.space_after = Pt(0)
    spacer.paragraph_format.line_spacing = Pt(3)
    run = spacer.add_run("\u00a0")
    set_font(run, size=1, color="FFFFFF")
    return doc.add_heading(text, level=level)


def configure(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.2

    for name, size, color, before, after in (
        ("Heading 1", 16, "2E74B5", 18, 10),
        ("Heading 2", 13, "2E74B5", 14, 7),
        ("Heading 3", 12, "1F4D78", 10, 5),
    ):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = header.add_run("NUTRISOFT  /  PRODUCT SYSTEM")
    set_font(run, size=8.5, color=TERTIARY, bold=True)
    add_page_number(section.footer.paragraphs[0])


def widths_for(headers):
    n = len(headers)
    if n == 2:
        return [2700, 6660]
    if n == 3:
        return [2300, 2500, 4560]
    if n == 4:
        return [1700, 2100, 2100, 3460]
    return [TABLE_WIDTH_DXA // n] * n


def add_table(doc, headers, rows, section_title):
    is_history = "Historial de decisiones" in section_title
    header_size = 8.0 if is_history else 8.7
    body_size = 7.6 if is_history else 8.5
    cell_vertical_margin = 45 if is_history else 80
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_geometry(table, widths_for(headers))
    prevent_split(table.rows[0])
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_shading(cell, "E8EEF5")
        set_cell_margins(cell, top=cell_vertical_margin, bottom=cell_vertical_margin)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        add_inline_runs(p, header, size=header_size, color=INK)
        for run in p.runs:
            run.bold = True
    is_palette = "Paleta oficial" in section_title
    for row_values in rows:
        row = table.add_row()
        prevent_split(row)
        for i, value in enumerate(row_values):
            cell = row.cells[i]
            set_cell_margins(cell, top=cell_vertical_margin, bottom=cell_vertical_margin)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            add_inline_runs(p, value, size=body_size)
        if is_palette and len(row_values) > 1:
            match = re.search(r"#[0-9A-Fa-f]{6}", row_values[1])
            if match:
                set_shading(row.cells[1], match.group(0)[1:].upper())
                lum = sum(int(match.group(0)[i:i+2], 16) for i in (1, 3, 5)) / 3
                for run in row.cells[1].paragraphs[0].runs:
                    run.font.color.rgb = RGBColor.from_string("FFFFFF" if lum < 135 else INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_callout(doc, text):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    set_table_geometry(table, [TABLE_WIDTH_DXA])
    row = table.rows[0]
    prevent_split(row)
    cell = row.cells[0]
    set_shading(cell, BRAND_SOFT)
    set_cell_margins(cell, top=130, bottom=130, start=180, end=180)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    add_inline_runs(p, text, size=10.3, color=BRAND_STRONG)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_code(doc, lines):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    set_table_geometry(table, [TABLE_WIDTH_DXA])
    cell = table.cell(0, 0)
    set_shading(cell, "F6F8FA")
    set_cell_margins(cell, top=120, bottom=120, start=160, end=160)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.05
    run = p.add_run("\n".join(lines))
    set_font(run, name="Consolas", size=8.7, color=INK)


def add_closing_rule(doc, text):
    """Render the final product principle as an intentional closing page."""
    label = doc.add_paragraph()
    label.paragraph_format.space_before = Pt(125)
    label.paragraph_format.space_after = Pt(12)
    run = label.add_run("CRITERIO DE CIERRE")
    set_font(run, size=9, color=BRAND_STRONG, bold=True)

    statement = doc.add_paragraph()
    statement.paragraph_format.space_after = Pt(20)
    statement.paragraph_format.line_spacing = 1.15
    run = statement.add_run(text)
    set_font(run, size=19, color=INK, bold=True)

    accent = doc.add_table(rows=1, cols=1)
    accent.autofit = False
    set_table_geometry(accent, [1700])
    cell = accent.cell(0, 0)
    set_shading(cell, BRAND)
    set_cell_margins(cell, top=18, bottom=18, start=0, end=0)
    cell.paragraphs[0].paragraph_format.space_after = Pt(0)


def build():
    markdown = SOURCE.read_text(encoding="utf-8")
    lines = markdown.splitlines()
    doc = Document()
    configure(doc)
    bullet_id, _ = add_numbering(doc)
    decimal_abstract_id = 60
    decimal_id = None
    next_decimal_id = 100
    props = doc.core_properties
    props.title = "NutriSoft — Documento Maestro de Estilo, Producto y Continuidad"
    props.subject = "Fuente de verdad de diseño, requisitos, seguridad y migración"
    props.author = "NutriSoft"
    props.keywords = "NutriSoft, design system, requisitos, seguridad, multi-tenant, continuidad"

    # Editorial cover (named override: branded left-aligned manual cover).
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(88)
    p.paragraph_format.space_after = Pt(12)
    r = p.add_run("NUTRISOFT")
    set_font(r, size=12, color=BRAND_STRONG, bold=True)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run("Documento Maestro")
    set_font(r, size=31, color=INK, bold=True)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(22)
    r = p.add_run("Estilo, producto, requisitos y continuidad")
    set_font(r, size=16, color=BRAND_STRONG)
    add_callout(doc, "Fuente de verdad para continuar, migrar, rediseñar o auditar NutriSoft sin perder decisiones importantes.")
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(30)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("Versión 1.1  ·  14 de agosto de 2026")
    set_font(r, size=10.5, color=SECONDARY, bold=True)
    p = doc.add_paragraph()
    r = p.add_run("Cierre funcional 2.2 verificado + backend seguro multi-tenant 2.1 verificado localmente")
    set_font(r, size=9.5, color=TERTIARY, italic=True)
    doc.add_page_break()

    current_section = ""
    i = 1  # skip Markdown H1
    while i < len(lines):
        line = lines[i].rstrip()
        stripped = line.strip()
        if not stripped:
            i += 1
            continue
        if stripped.startswith("**Versión:**") or stripped.startswith("**Fecha de corte:**") or stripped.startswith("**Estado del producto:**") or stripped.startswith("**Propósito:**"):
            i += 1
            continue
        if stripped.startswith("**Regla final:**"):
            closing_text = re.sub(r"^\*\*Regla final:\*\*\s*", "", stripped)
            add_closing_rule(doc, closing_text)
            i += 1
            continue
        if stripped == "---":
            i += 1
            continue
        if stripped.startswith("## "):
            current_section = stripped[3:]
            add_section_heading(doc, current_section, level=1)
            i += 1
            continue
        if stripped.startswith("### "):
            current_section = stripped[4:]
            add_section_heading(doc, current_section, level=2)
            i += 1
            continue
        if stripped.startswith("#### "):
            add_section_heading(doc, stripped[5:], level=3)
            i += 1
            continue
        if stripped.startswith("> "):
            add_callout(doc, stripped[2:])
            i += 1
            continue
        if stripped.startswith("```"):
            code_lines = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code_lines.append(lines[i])
                i += 1
            add_code(doc, code_lines)
            i += 1
            continue
        if stripped.startswith("|"):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i].strip())
                i += 1
            parsed = [[cell.strip() for cell in row.strip("|").split("|")] for row in table_lines]
            headers = parsed[0]
            rows = [row for row in parsed[2:] if not all(re.fullmatch(r":?-+:?", c) for c in row)]
            add_table(doc, headers, rows, current_section)
            continue
        if re.match(r"^- ", stripped):
            p = doc.add_paragraph()
            set_num(p, bullet_id)
            add_inline_runs(p, stripped[2:])
            i += 1
            continue
        m = re.match(r"^(\d+)\.\s+(.*)$", stripped)
        if m:
            # Markdown lists that start at 1 are independent sequences. Giving
            # each one its own Word numId prevents numbering from continuing
            # across unrelated sections.
            if m.group(1) == "1" or decimal_id is None:
                decimal_id = add_numbering_instance(doc, decimal_abstract_id, next_decimal_id)
                next_decimal_id += 1
            p = doc.add_paragraph()
            set_num(p, decimal_id)
            add_inline_runs(p, m.group(2))
            i += 1
            continue

        # Merge wrapped prose lines until the next structural line.
        paragraph_lines = [stripped]
        i += 1
        while i < len(lines):
            nxt = lines[i].strip()
            if not nxt:
                i += 1
                break
            if nxt.startswith(("#", "|", ">", "```", "- ")) or re.match(r"^\d+\.\s+", nxt):
                break
            paragraph_lines.append(nxt)
            i += 1
        p = doc.add_paragraph()
        add_inline_runs(p, " ".join(paragraph_lines))

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
